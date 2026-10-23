import json
import re
from collections.abc import AsyncIterator, Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from langchain_core.callbacks import AsyncCallbackManagerForLLMRun, CallbackManagerForLLMRun
from langchain_core.language_models import BaseChatModel
from langchain_core.messages import AIMessage, AIMessageChunk, BaseMessage
from langchain_core.outputs import ChatGeneration, ChatGenerationChunk, ChatResult
from pydantic import Field
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_PATH = Path(__file__).resolve().parents[2] / "scripts" / "fonts" / "DejaVuSans.ttf"
_FONT_NAME = "DejaVuSans"


def make_pdf(path: Path, pages: list[str]) -> Path:
    """PDF with one page per item; an empty string gives a page without text."""
    if _FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(_FONT_NAME, str(FONT_PATH)))
    c = canvas.Canvas(str(path))
    for text in pages:
        if text:
            c.setFont(_FONT_NAME, 10)
            y = 800
            line = ""
            for word in text.split(" "):
                candidate = f"{line} {word}".strip()
                if len(candidate) > 90:
                    c.drawString(40, y, line)
                    y -= 14
                    line = word
                else:
                    line = candidate
            if line:
                c.drawString(40, y, line)
        c.showPage()
    c.save()
    return path


def make_docx(path: Path, blocks: list[tuple[str, str]]) -> Path:
    """Blocks: ("h1" | "h2" | "p", text) or ("table", "a|b;c|d")."""
    doc = DocxDocument()
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "p":
            doc.add_paragraph(text)
        elif kind == "table":
            rows = [r.split("|") for r in text.split(";")]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
        else:
            raise ValueError(f"unknown block kind: {kind}")
    doc.save(str(path))
    return path


# --- Scripted chat model for agent tests ---


@dataclass
class RaiseAfter:
    """Scripted step: stream `text`, then raise `exc` (simulates a model error mid-stream)."""

    text: str
    exc: Exception


def _split_text(text: str) -> list[str]:
    """Split into word-ish pieces; at least 2 pieces when the text has 2+ characters."""
    parts = re.findall(r"\S+\s*|\s+", text)
    if len(parts) < 2 and len(text) >= 2:
        mid = len(text) // 2
        parts = [text[:mid], text[mid:]]
    return parts


class ScriptedChatModel(BaseChatModel):
    """Returns scripted responses in order; records every input message list in `calls`."""

    responses: list[Any]
    calls: list[list[BaseMessage]] = Field(default_factory=list)
    index: int = 0

    @property
    def _llm_type(self) -> str:
        return "scripted-fake"

    def bind_tools(self, tools: Sequence[Any], **kwargs: Any) -> "ScriptedChatModel":
        return self

    def _next(self, messages: list[BaseMessage]) -> AIMessage | RaiseAfter:
        self.calls.append(list(messages))
        if self.index >= len(self.responses):
            raise RuntimeError(
                f"ScriptedChatModel ran out of responses: {len(self.responses)} scripted, call #{self.index + 1}"
            )
        response = self.responses[self.index]
        self.index += 1
        if not isinstance(response, (AIMessage, RaiseAfter)):
            raise TypeError(f"ScriptedChatModel: unsupported scripted item {response!r}")
        return response

    def _generate(
        self,
        messages: list[BaseMessage],
        stop: list[str] | None = None,
        run_manager: CallbackManagerForLLMRun | None = None,
        **kwargs: Any,
    ) -> ChatResult:
        response = self._next(messages)
        if isinstance(response, RaiseAfter):
            raise response.exc
        return ChatResult(generations=[ChatGeneration(message=response.model_copy())])

    def _chunks(self, response: AIMessage | RaiseAfter) -> Iterator[ChatGenerationChunk]:
        text = response.text if isinstance(response, RaiseAfter) else response.content
        if not isinstance(text, str):
            raise TypeError("ScriptedChatModel supports only string content")
        for part in _split_text(text):
            yield ChatGenerationChunk(message=AIMessageChunk(content=part))
        if isinstance(response, RaiseAfter):
            raise response.exc
        if response.tool_calls:
            yield ChatGenerationChunk(
                message=AIMessageChunk(
                    content="",
                    tool_call_chunks=[
                        {"name": tc["name"], "args": json.dumps(tc["args"], ensure_ascii=False), "id": tc["id"], "index": i}
                        for i, tc in enumerate(response.tool_calls)
                    ],
                )
            )

    def _stream(
        self,
        messages: list[BaseMessage],
        stop: list[str] | None = None,
        run_manager: CallbackManagerForLLMRun | None = None,
        **kwargs: Any,
    ) -> Iterator[ChatGenerationChunk]:
        for chunk in self._chunks(self._next(messages)):
            if run_manager:
                run_manager.on_llm_new_token(chunk.text, chunk=chunk)
            yield chunk

    async def _astream(
        self,
        messages: list[BaseMessage],
        stop: list[str] | None = None,
        run_manager: AsyncCallbackManagerForLLMRun | None = None,
        **kwargs: Any,
    ) -> AsyncIterator[ChatGenerationChunk]:
        for chunk in self._chunks(self._next(messages)):
            if run_manager:
                await run_manager.on_llm_new_token(chunk.text, chunk=chunk)
            yield chunk